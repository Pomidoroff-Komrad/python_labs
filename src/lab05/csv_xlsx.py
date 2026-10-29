from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from src.lab05.json_csv import read_csv


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """Переносит CSV на лист Sheet1 и подбирает ширину колонок."""
    out = Path(xlsx_path)
    if out.suffix.lower() != '.xlsx':
        raise ValueError('Нужен выходной файл XLSX')
    rows = read_csv(csv_path)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    for row in rows:
        ws.append(row)
        # CSV хранит текст: строки с = тоже остаются текстом, а не формулами.
        for cell in ws[ws.max_row]:
            cell.data_type = 's'

    for i in range(len(rows[0])):
        width = 8
        for row in rows:
            width = max(width, len(row[i]) + 2)
        column = get_column_letter(i + 1)
        ws.column_dimensions[column].width = width

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    wb.close()


if __name__ == '__main__':
    csv_to_xlsx('data/samples/people.csv', 'data/out/people.xlsx')
    csv_to_xlsx('data/samples/cities.csv', 'data/out/cities.xlsx')
